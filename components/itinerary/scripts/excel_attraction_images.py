"""Extract per-row images embedded in attractions.xlsx (Excel rich-data IMAGE cells)."""

from __future__ import annotations

import base64
import posixpath
import re
import shutil
import urllib.parse
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[3]
BACKEND_DIR = REPO_ROOT / "components" / "shared" / "backend"
DEFAULT_XLSX = Path(__file__).resolve().parents[1] / "data" / "attractions.xlsx"
DEFAULT_OUT_DIR = Path(__file__).resolve().parents[1] / "backend" / "static" / "attractions"
PUBLIC_URL_PREFIX = "/api/media/attractions"
MANUAL_PHOTO_DIR = Path(__file__).resolve().parents[1] / "data" / "manual_attraction_photos"
# Applied after Excel extract (which wipes static/attractions).
MANUAL_PHOTO_BY_NAME = {
    "baron's cap or thoppigala": "042-baron-s-cap-or-thoppigala.jpeg",
    "st. clair's falls": "040-st-clair-s-falls.jpeg",
    "little adam's peak": "048-little-adam-s-peak.jpeg",
    "lover's leap waterfall": "lovers-leap-waterfall.jpeg",
    "ravana's cave": "ravanas-cave.jpeg",
}

NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}


def _col_letters_to_index(col: str) -> int:
    n = 0
    for ch in col.upper():
        n = n * 26 + (ord(ch) - ord("A") + 1)
    return n - 1


def _parse_cell_ref(ref: str) -> tuple[int, int]:
    m = re.match(r"^([A-Z]+)(\d+)$", ref.upper())
    if not m:
        raise ValueError(f"Bad cell ref: {ref}")
    return int(m.group(2)), _col_letters_to_index(m.group(1))


def _load_rels(zf: zipfile.ZipFile, rels_path: str) -> dict[str, str]:
    if rels_path not in zf.namelist():
        return {}
    root = ET.fromstring(zf.read(rels_path))
    mapping: dict[str, str] = {}
    for rel in root:
        rid = rel.attrib.get("Id")
        target = rel.attrib.get("Target")
        if rid and target:
            mapping[rid] = target.replace("\\", "/")
    return mapping


def _resolve_zip_path(base_part: str, target: str) -> str:
    """Join a package-relative target and normalize '..' for zip namelist lookups."""
    if target.startswith("/"):
        return target.lstrip("/")
    base_dir = posixpath.dirname(base_part.replace("\\", "/"))
    return posixpath.normpath(posixpath.join(base_dir, target.replace("\\", "/")))


def _rich_data_media_by_vm(zf: zipfile.ZipFile) -> dict[int, str]:
    """Map 1-based vm index -> zip path under xl/media/..."""
    rels = _load_rels(zf, "xl/richData/_rels/richValueRel.xml.rels")
    if not rels:
        return {}

    # Prefer ordered richValueRel.xml rId list when present
    ordered_rids: list[str] = []
    if "xl/richData/richValueRel.xml" in zf.namelist():
        root = ET.fromstring(zf.read("xl/richData/richValueRel.xml"))
        for el in root.iter():
            rid = el.attrib.get(f"{{{NS['r']}}}id") or el.attrib.get("r:id")
            if rid:
                ordered_rids.append(rid)

    by_vm: dict[int, str] = {}
    if ordered_rids:
        for idx, rid in enumerate(ordered_rids, start=1):
            target = rels.get(rid)
            if not target:
                continue
            by_vm[idx] = _resolve_zip_path("xl/richData/richValueRel.xml", target)
        return by_vm

    # Fallback: rIdN -> media (works for this workbook)
    for rid, target in rels.items():
        m = re.match(r"rId(\d+)$", rid)
        if not m:
            continue
        by_vm[int(m.group(1))] = _resolve_zip_path("xl/richData/richValueRel.xml", target)
    return by_vm


def _image_column_index(zf: zipfile.ZipFile, sheet_path: str) -> int | None:
    root = ET.fromstring(zf.read(sheet_path))
    header_row = root.find("main:sheetData/main:row[@r='1']", NS)
    if header_row is None:
        return None
    for cell in header_row.findall("main:c", NS):
        ref = cell.attrib.get("r", "")
        row, col = _parse_cell_ref(ref)
        if row != 1:
            continue
        # shared string or inline
        v = cell.find("main:v", NS)
        is_ = cell.find("main:is/main:t", NS)
        text = ""
        if cell.attrib.get("t") == "s" and v is not None and v.text is not None:
            text = _shared_string(zf, int(v.text))
        elif is_ is not None and is_.text:
            text = is_.text
        elif v is not None and v.text:
            text = v.text
        if text.strip().lower() == "image":
            return col
    return 5  # column F in the known workbook


def _shared_string(zf: zipfile.ZipFile, idx: int) -> str:
    root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
    items = root.findall("main:si", NS)
    if idx < 0 or idx >= len(items):
        return ""
    texts = [t.text or "" for t in items[idx].iter("{%s}t" % NS["main"])]
    return "".join(texts)


def _vm_by_excel_row(zf: zipfile.ZipFile, sheet_path: str, image_col: int) -> dict[int, int]:
    root = ET.fromstring(zf.read(sheet_path))
    mapping: dict[int, int] = {}
    for cell in root.findall("main:sheetData/main:row/main:c", NS):
        ref = cell.attrib.get("r")
        vm = cell.attrib.get("vm")
        if not ref or not vm:
            continue
        row, col = _parse_cell_ref(ref)
        if col != image_col or row < 2:
            continue
        mapping[row] = int(vm)
    return mapping


def _drawing_media_by_row(zf: zipfile.ZipFile, sheet_path: str) -> dict[int, str]:
    sheet_rels_path = (
        Path(sheet_path).parent.as_posix()
        + "/_rels/"
        + Path(sheet_path).name
        + ".rels"
    )
    sheet_rels = _load_rels(zf, sheet_rels_path)
    drawing_targets = [
        t for t in sheet_rels.values() if "drawing" in t.lower()
    ]
    result: dict[int, str] = {}
    for target in drawing_targets:
        drawing_path = _resolve_zip_path(sheet_path, target)
        if drawing_path not in zf.namelist():
            continue
        drawing_rels = _load_rels(
            zf,
            Path(drawing_path).parent.as_posix()
            + "/_rels/"
            + Path(drawing_path).name
            + ".rels",
        )
        root = ET.fromstring(zf.read(drawing_path))
        for anchor in list(root):
            frm = anchor.find("xdr:from", NS)
            if frm is None:
                continue
            col_el = frm.find("xdr:col", NS)
            row_el = frm.find("xdr:row", NS)
            if col_el is None or row_el is None:
                continue
            excel_row = int(row_el.text) + 1  # 0-based in drawing
            blip = anchor.find(".//a:blip", NS)
            if blip is None:
                continue
            rid = blip.attrib.get(f"{{{NS['r']}}}embed") or blip.attrib.get("r:embed")
            if not rid or rid not in drawing_rels:
                continue
            media = _resolve_zip_path(drawing_path, drawing_rels[rid])
            result[excel_row] = media
    return result


def map_excel_row_to_media(xlsx_path: Path) -> dict[int, str]:
    """Return excel_row -> zip media path (e.g. xl/media/image1.jpeg)."""
    with zipfile.ZipFile(xlsx_path, "r") as zf:
        sheet_path = "xl/worksheets/sheet1.xml"
        if sheet_path not in zf.namelist():
            # fallback first worksheet
            sheets = [n for n in zf.namelist() if n.startswith("xl/worksheets/sheet")]
            if not sheets:
                return {}
            sheet_path = sorted(sheets)[0]

        image_col = _image_column_index(zf, sheet_path)
        if image_col is None:
            image_col = 5

        rich_by_vm = _rich_data_media_by_vm(zf)
        vm_by_row = _vm_by_excel_row(zf, sheet_path, image_col)
        drawing_by_row = _drawing_media_by_row(zf, sheet_path)

        mapped: dict[int, str] = {}
        for row, vm in vm_by_row.items():
            media = rich_by_vm.get(vm)
            if media:
                mapped[row] = media
        for row, media in drawing_by_row.items():
            mapped.setdefault(row, media)
        return mapped


def _slugify(name: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", (name or "").strip().lower())
    return slug.strip("-")[:60] or "attraction"


def extract_attraction_images(
    xlsx_path: Path | None = None,
    out_dir: Path | None = None,
    attraction_names_by_excel_row: dict[int, str] | None = None,
) -> dict[int, str]:
    """
    Extract embedded images into out_dir.

    Returns excel_row -> public URL path (/api/media/attractions/...).
    """
    xlsx_path = Path(xlsx_path or DEFAULT_XLSX)
    out_dir = Path(out_dir or DEFAULT_OUT_DIR)
    if out_dir.exists():
        shutil.rmtree(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    media_by_row = map_excel_row_to_media(xlsx_path)
    urls: dict[int, str] = {}

    with zipfile.ZipFile(xlsx_path, "r") as zf:
        for excel_row, media_path in sorted(media_by_row.items()):
            if media_path not in zf.namelist():
                continue
            ext = Path(media_path).suffix.lower() or ".jpeg"
            name = (attraction_names_by_excel_row or {}).get(excel_row, "")
            slug = _slugify(name)
            filename = f"{excel_row:03d}-{slug}{ext}" if slug != "attraction" else f"{excel_row:03d}{ext}"
            dest = out_dir / filename
            dest.write_bytes(zf.read(media_path))
            urls[excel_row] = f"{PUBLIC_URL_PREFIX}/{filename}"

    return urls


def normalize_image_url(value) -> str | None:
    """Turn an Excel Image cell value into a usable img src URL."""
    if value is None:
        return None
    text = str(value).strip()
    if not text or text.lower() in ("nan", "none", "#value!"):
        return None
    if text.startswith("data:image"):
        return text
    if not text.startswith("http"):
        return None

    try:
        parsed = urllib.parse.urlparse(text)
    except Exception:
        return None

    host = (parsed.netloc or "").lower()
    if "google." in host and ("/imgres" in parsed.path or "imgurl=" in text):
        qs = urllib.parse.parse_qs(parsed.query)
        imgurl = qs.get("imgurl", [None])[0]
        if imgurl:
            return urllib.parse.unquote(imgurl)
        return None

    if "google." in host and "/search" in parsed.path:
        return None

    return text


def _save_base64_image(data_uri: str, excel_row: int, name: str, out_dir: Path) -> str | None:
    """Decode data:image/...;base64,... and save under static/attractions."""
    match = re.match(r"^data:image/([\w+.-]+);base64,(.+)$", data_uri.strip(), re.DOTALL)
    if not match:
        return None
    ext_raw = match.group(1).lower().replace("jpeg", "jpg")
    ext = f".{ext_raw}" if ext_raw else ".jpg"
    try:
        raw = base64.b64decode(match.group(2), validate=False)
    except Exception:
        return None
    if not raw:
        return None

    slug = _slugify(name)
    filename = f"{excel_row:03d}-{slug}{ext}" if slug != "attraction" else f"{excel_row:03d}{ext}"
    out_dir.mkdir(parents=True, exist_ok=True)
    dest = out_dir / filename
    dest.write_bytes(raw)
    return f"{PUBLIC_URL_PREFIX}/{filename}"


def merge_cell_image_urls(
    xlsx_path: Path | None,
    attraction_names: list[str],
    embedded_urls: list[str | None],
    out_dir: Path | None = None,
) -> list[str | None]:
    """
    Fill gaps in embedded_urls from Excel Image column:
    - plain https://... URLs pasted in cells
    - hyperlinks on the cell
    - data:image base64 blobs (saved to static/attractions)
    """
    try:
        import openpyxl
    except ImportError:
        return list(embedded_urls)

    xlsx_path = Path(xlsx_path or DEFAULT_XLSX)
    out_dir = Path(out_dir or DEFAULT_OUT_DIR)
    merged = list(embedded_urls)
    if len(merged) != len(attraction_names):
        merged = (merged + [None] * len(attraction_names))[: len(attraction_names)]

    wb = openpyxl.load_workbook(xlsx_path, data_only=False)
    ws = wb["Attractions"] if "Attractions" in wb.sheetnames else wb.active
    headers = [c.value for c in ws[1]]
    if "Image" not in headers:
        return merged
    img_col = headers.index("Image") + 1

    for i, name in enumerate(attraction_names):
        if merged[i]:
            continue
        excel_row = i + 2
        cell = ws.cell(row=excel_row, column=img_col)
        raw = None
        if cell.hyperlink and cell.hyperlink.target:
            raw = str(cell.hyperlink.target).strip()
        elif cell.value is not None:
            raw = str(cell.value).strip()

        if not raw or raw.lower() in ("#value!", "none", "nan"):
            continue

        if raw.startswith("data:image"):
            saved = _save_base64_image(raw, excel_row, name, out_dir)
            if saved:
                merged[i] = saved
            continue

        url = normalize_image_url(raw)
        if url:
            merged[i] = url

    return merged


def apply_manual_attraction_photos(
    attraction_names: list[str],
    urls: list[str | None],
    out_dir: Path | None = None,
) -> list[str | None]:
    """Copy curator photos over Excel embeds / broken Google thumbs."""
    if not MANUAL_PHOTO_DIR.is_dir():
        return urls
    dest_dir = Path(out_dir or DEFAULT_OUT_DIR)
    dest_dir.mkdir(parents=True, exist_ok=True)
    merged = list(urls)
    if len(merged) != len(attraction_names):
        merged = (merged + [None] * len(attraction_names))[: len(attraction_names)]
    for i, name in enumerate(attraction_names):
        filename = MANUAL_PHOTO_BY_NAME.get((name or "").strip().lower())
        if not filename:
            continue
        src = MANUAL_PHOTO_DIR / filename
        if not src.is_file():
            continue
        shutil.copy2(src, dest_dir / filename)
        merged[i] = f"{PUBLIC_URL_PREFIX}/{filename}"
    return merged


def attraction_image_urls_in_order(
    xlsx_path: Path | None = None,
    out_dir: Path | None = None,
    attraction_names: list[str] | None = None,
) -> list[str | None]:
    """
    Extract images and return a list aligned to pandas row order
    (excel row 2 == index 0).

    Uses embedded Excel images first, then cell URLs / hyperlinks / base64.
    """
    names = attraction_names or []
    names_by_row = {i + 2: name for i, name in enumerate(names)}
    urls_by_row = extract_attraction_images(
        xlsx_path=xlsx_path,
        out_dir=out_dir,
        attraction_names_by_excel_row=names_by_row,
    )
    embedded = [urls_by_row.get(i + 2) for i in range(len(names))]
    merged = merge_cell_image_urls(
        xlsx_path=xlsx_path,
        attraction_names=names,
        embedded_urls=embedded,
        out_dir=out_dir,
    )
    return apply_manual_attraction_photos(names, merged, out_dir=out_dir)


if __name__ == "__main__":
    import pandas as pd

    xlsx = DEFAULT_XLSX
    df = pd.read_excel(xlsx, sheet_name="Attractions")
    names = [str(n) if n is not None else "" for n in df["Attraction Name"].tolist()]
    urls = attraction_image_urls_in_order(attraction_names=names)
    found = sum(1 for u in urls if u)
    print(f"Extracted {found}/{len(urls)} attraction images -> {DEFAULT_OUT_DIR}")
    for name, url in list(zip(names, urls))[:8]:
        print(f"  {name}: {url}")
